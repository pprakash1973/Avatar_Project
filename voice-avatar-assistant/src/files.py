"""Read an uploaded file (PDF, Word, or image) with Claude and summarize it."""

import base64
import io
import json
from dataclasses import dataclass
from pathlib import PurePath

from docx import Document

from claude import HELPER_MODEL, HelperUsage, get_client

MAX_FILE_BYTES = 1_000_000

_KINDS_BY_EXTENSION = {
    ".pdf": ("pdf", "application/pdf"),
    ".docx": ("docx", None),
    ".jpg": ("image", "image/jpeg"),
    ".jpeg": ("image", "image/jpeg"),
    ".png": ("image", "image/png"),
}

_DIGEST_SCHEMA = {
    "type": "object",
    "properties": {
        "title": {"type": "string"},
        "summary": {"type": "string"},
        "key_points": {"type": "array", "items": {"type": "string"}},
        "notes": {"type": "string"},
    },
    "required": ["title", "summary", "key_points", "notes"],
    "additionalProperties": False,
}


class UnsupportedFileError(Exception):
    pass


@dataclass
class FileDigest:
    title: str
    summary: str
    key_points: list[str]
    notes: str


def file_kind(name: str) -> tuple[str, str | None]:
    kind = _KINDS_BY_EXTENSION.get(PurePath(name).suffix.lower())
    if kind is None:
        raise UnsupportedFileError(name)
    return kind


def _docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def _file_block(name: str, data: bytes) -> dict:
    kind, media_type = file_kind(name)
    if kind == "docx":
        text = _docx_text(data)
        if not text.strip():
            raise UnsupportedFileError(name)
        return {
            "type": "text",
            "text": f"<document name={name!r}>\n{text}\n</document>",
        }
    encoded = base64.standard_b64encode(data).decode("ascii")
    source = {"type": "base64", "media_type": media_type, "data": encoded}
    return {"type": "document" if kind == "pdf" else "image", "source": source}


async def summarize_file(
    name: str, data: bytes, language: str, usage: HelperUsage
) -> FileDigest:
    response = await get_client().messages.create(
        model=HELPER_MODEL,
        max_tokens=4000,
        output_config={
            "effort": "low",
            "format": {"type": "json_schema", "schema": _DIGEST_SCHEMA},
        },
        messages=[
            {
                "role": "user",
                "content": [
                    _file_block(name, data),
                    {
                        "type": "text",
                        "text": (
                            f"The user uploaded the file {name!r}. Write all fields in "
                            f"{language}, as plain text without markdown.\n"
                            "- title: a short title for the file\n"
                            "- summary: two or three sentences summarizing it, suitable "
                            "for reading aloud\n"
                            "- key_points: three to six short key points\n"
                            "- notes: detailed notes (up to about 800 words) with the "
                            "important facts, figures, names and dates, so follow-up "
                            "questions can be answered without the file"
                        ),
                    },
                ],
            }
        ],
    )
    usage.add(response.usage)
    text = next(b.text for b in response.content if b.type == "text")
    return FileDigest(**json.loads(text))
